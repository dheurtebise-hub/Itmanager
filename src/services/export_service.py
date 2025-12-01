"""
Export des tickets en CSV et Excel
"""

import csv
import io
from datetime import datetime
from typing import List, Dict, Any

class ExportService:
    FIELDS = [
        ('id', 'ID'), ('subject', 'Sujet'), ('sender_email', 'Expéditeur'),
        ('category', 'Catégorie'), ('priority', 'Priorité'), ('status', 'Statut'),
        ('received_date', 'Date réception'), ('resolved_at', 'Date résolution'),
        ('resolution', 'Résolution'), ('summary', 'Résumé')
    ]

    def export_csv(self, tickets: List[Dict[str, Any]]) -> str:
        output = io.StringIO()
        writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL)
        writer.writerow([label for _, label in self.FIELDS])

        for ticket in tickets:
            writer.writerow([self._format(ticket.get(f, '')) for f, _ in self.FIELDS])

        return output.getvalue()

    def export_excel(self, tickets: List[Dict[str, Any]]) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Tickets"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")

        for col, (_, label) in enumerate(self.FIELDS, 1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = header_font
            cell.fill = header_fill

        for row_num, ticket in enumerate(tickets, 2):
            for col, (field, _) in enumerate(self.FIELDS, 1):
                ws.cell(row=row_num, column=col, value=self._format(ticket.get(field, '')))

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _format(self, value: Any) -> str:
        if value is None:
            return ''
        if isinstance(value, datetime):
            return value.strftime('%d/%m/%Y %H:%M')
        return str(value)


export_service = ExportService()
